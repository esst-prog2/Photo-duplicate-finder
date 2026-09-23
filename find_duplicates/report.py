from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from PIL import Image

KEEP_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def _keep_metric(path: Path) -> tuple[int, int]:
    file_size = path.stat().st_size
    with Image.open(path) as img:
        resolution = img.width * img.height
    return (file_size, resolution)


def choose_keep(group: list[Path]) -> Path:
    return max(group, key=_keep_metric)


@dataclass
class DuplicateGroup:
    members: list[Path]
    similarity_score: str
    keep: Path


def write_report(groups: list[DuplicateGroup], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active

    row = 1
    for group in groups:
        sheet.cell(row=row, column=1, value="File")
        sheet.cell(row=row, column=2, value="Similarity Score")
        row += 1
        for member in group.members:
            file_cell = sheet.cell(row=row, column=1, value=member.name)
            score_cell = sheet.cell(row=row, column=2, value=group.similarity_score)
            if member == group.keep:
                file_cell.fill = KEEP_FILL
                score_cell.fill = KEEP_FILL
            row += 1
        row += 1

    workbook.save(output_path)
