from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from find_duplicates.report import KEEP_FILL, DuplicateGroup, choose_keep, write_report

FIXTURES = Path(__file__).parent / "fixtures"


def test_largest_file_size_wins_even_with_lower_resolution():
    base = FIXTURES / "base.png"
    resized = FIXTURES / "base_resized.png"

    # base_resized.png has a lower resolution than base.png but a larger file
    # size (PNG compression of the downscaled image happens to be less
    # efficient here) -- file size takes priority over resolution.
    assert resized.stat().st_size > base.stat().st_size
    with Image.open(base) as base_img, Image.open(resized) as resized_img:
        assert (resized_img.width * resized_img.height) < (base_img.width * base_img.height)

    assert choose_keep([base, resized]) == resized


def test_tie_break_by_resolution_when_file_sizes_are_equal(tmp_path):
    small = Image.new("RGB", (50, 50), color=(10, 20, 30))
    large = Image.new("RGB", (200, 200), color=(10, 20, 30))
    small_path = tmp_path / "small.png"
    large_path = tmp_path / "large.png"
    small.save(small_path)
    large.save(large_path)

    size_diff = large_path.stat().st_size - small_path.stat().st_size
    assert size_diff >= 0
    with open(small_path, "ab") as f:
        f.write(b"\x00" * size_diff)
    assert small_path.stat().st_size == large_path.stat().st_size

    assert choose_keep([small_path, large_path]) == large_path


def test_write_report_writes_one_table_per_group_with_blank_row_between(tmp_path):
    base = FIXTURES / "base.png"
    resized = FIXTURES / "base_resized.png"
    unrelated = FIXTURES / "unrelated.png"
    groups = [
        DuplicateGroup(members=[base, resized], similarity_score="Very Similar", keep=resized),
        DuplicateGroup(members=[base, unrelated], similarity_score="Similar", keep=base),
    ]
    output_path = tmp_path / "duplicates.xlsx"

    write_report(groups, output_path)

    sheet = load_workbook(output_path).active
    rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows(max_col=2)]

    assert rows == [
        ("File", "Similarity Score"),
        ("base.png", "Very Similar"),
        ("base_resized.png", "Very Similar"),
        (None, None),
        ("File", "Similarity Score"),
        ("base.png", "Similar"),
        ("unrelated.png", "Similar"),
    ]


def test_write_report_highlights_the_keep_row(tmp_path):
    base = FIXTURES / "base.png"
    resized = FIXTURES / "base_resized.png"
    groups = [DuplicateGroup(members=[base, resized], similarity_score="Very Similar", keep=resized)]
    output_path = tmp_path / "duplicates.xlsx"

    write_report(groups, output_path)

    sheet = load_workbook(output_path).active
    # Row 2 is base.png (not kept), row 3 is base_resized.png (kept).
    assert sheet.cell(row=2, column=1).fill.start_color.rgb != KEEP_FILL.start_color.rgb
    assert sheet.cell(row=3, column=1).fill.start_color.rgb == KEEP_FILL.start_color.rgb
    assert sheet.cell(row=3, column=2).fill.start_color.rgb == KEEP_FILL.start_color.rgb


def test_write_report_with_no_groups_creates_valid_empty_workbook(tmp_path):
    output_path = tmp_path / "duplicates.xlsx"

    write_report([], output_path)

    assert output_path.exists()
    sheet = load_workbook(output_path).active
    assert sheet.max_row == 1
    assert sheet.cell(row=1, column=1).value is None
