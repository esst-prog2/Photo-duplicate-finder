import csv
import shutil
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from find_duplicates.cli import main

BENCHMARK = Path(__file__).parent.parent / "benchmark"

# Expected values were decided before the program ran on this data and are
# logged in PLANNING_LOG.md (2026-10-09). They come from planted data: 20
# hand-picked distinct photos, each given 3 scripted copies by
# benchmark/generate.py. They are not derived from the program's output.
EXPECTED_GROUPS = 20
EXPECTED_GROUP_SIZE = 4
EXPECTED_CONTAMINATED_GROUPS = 0


def _read_true_groups() -> dict[str, str]:
    with open(BENCHMARK / "manifest.csv", newline="", encoding="utf-8") as manifest:
        return {row["file"]: row["group"] for row in csv.DictReader(manifest)}


def _read_reported_groups(report_path: Path) -> list[list[str]]:
    sheet = load_workbook(report_path).active
    groups: list[list[str]] = []
    current: list[str] | None = None
    for file_cell, _score in sheet.iter_rows(min_col=1, max_col=2, values_only=True):
        if file_cell is None:
            current = None
        elif file_cell == "File":
            current = []
            groups.append(current)
        else:
            current.append(file_cell)
    return groups


@pytest.fixture(scope="module")
def reported_groups(tmp_path_factory) -> list[list[str]]:
    workspace = tmp_path_factory.mktemp("benchmark")
    folder = workspace / "data"
    shutil.copytree(BENCHMARK / "data", folder)

    result = CliRunner().invoke(main, [str(folder)])

    assert result.exit_code == 0, result.output
    report_paths = list(workspace.glob("duplicates_*.xlsx"))
    assert len(report_paths) == 1
    return _read_reported_groups(report_paths[0])


def test_benchmark_reports_the_expected_number_of_groups(reported_groups):
    assert len(reported_groups) == EXPECTED_GROUPS


def test_benchmark_every_group_has_the_expected_size(reported_groups):
    sizes = sorted(len(group) for group in reported_groups)

    assert sizes == [EXPECTED_GROUP_SIZE] * EXPECTED_GROUPS


def test_benchmark_has_no_contaminated_groups(reported_groups):
    true_groups = _read_true_groups()

    contaminated = [
        group for group in reported_groups if len({true_groups[name] for name in group}) > 1
    ]

    assert len(contaminated) == EXPECTED_CONTAMINATED_GROUPS, contaminated
