from click.testing import CliRunner
from openpyxl import load_workbook

from find_duplicates.cli import main


def test_empty_folder_exits_cleanly_with_zero_counts(tmp_path):
    empty_folder = tmp_path / "photos"
    empty_folder.mkdir()

    runner = CliRunner()
    result = runner.invoke(main, [str(empty_folder)])

    assert result.exit_code == 0
    assert result.output.strip() == "0 exact duplicates found, 0 near-duplicate groups found."

    report_path = tmp_path / "duplicates.xlsx"
    assert report_path.exists()
    sheet = load_workbook(report_path).active
    assert sheet.max_row == 1
    assert sheet.cell(row=1, column=1).value is None


def test_folder_with_no_matching_images_exits_cleanly_with_zero_counts(tmp_path):
    folder = tmp_path / "photos"
    folder.mkdir()
    (folder / "notes.txt").write_text("not an image")

    runner = CliRunner()
    result = runner.invoke(main, [str(folder)])

    assert result.exit_code == 0
    assert result.output.strip() == "0 exact duplicates found, 0 near-duplicate groups found."

    report_path = tmp_path / "duplicates.xlsx"
    assert report_path.exists()
    sheet = load_workbook(report_path).active
    assert sheet.max_row == 1
    assert sheet.cell(row=1, column=1).value is None
